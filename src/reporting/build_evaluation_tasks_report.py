"""Build a task execution report from saved results; no fitting, scoring or final project report."""
from pathlib import Path
import sys,json,csv,html,statistics
from collections import Counter
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT))
from src.training.nucle_final_selection import OUT,read,dump,LABELS,FULL,MASK,NONLEX,sha

NAMES={'Majority':'Majority','POS':'POS rule'}
def fmt(v):return 'N/A' if v is None else f'{v:.4f}'
def ci(v):return 'N/A' if v is None else f'[{v[0]:.4f}; {v[1]:.4f}]'
def write_csv(path,items):
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(items[0]));w.writeheader();w.writerows(items)

class Report:
    def __init__(self):
        self.doc=Document();self.md=[];self.eq=0
        sec=self.doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.7)
        style=self.doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(10)
        style.paragraph_format.space_after=Pt(6)
        footer=sec.footer.paragraphs[0];footer.text='NUCLE • Báo cáo task • 23/10/2026 • '
        fld=OxmlElement('w:fldSimple');fld.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr','PAGE');footer._p.append(fld)
    def h(self,text,level=1):self.doc.add_heading(text,level=level);self.md+=['#'*level+' '+text,'']
    def p(self,text):self.doc.add_paragraph(text);self.md+=[text,'']
    def table(self,headers,values):
        t=self.doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
        for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
        for row in values:
            for c,v in zip(t.add_row().cells,row):c.text=str(v)
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:run.font.size=Pt(8)
        self.md+=['| '+' | '.join(headers)+' |','|'+'|'.join(['---']*len(headers))+'|']
        self.md+=['| '+' | '.join(str(v).replace('|','/') for v in row)+' |' for row in values];self.md+=['']
    def math(self,text):
        self.eq+=1;p=self.doc.add_paragraph();om=OxmlElement('m:oMathPara');math=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=f'({self.eq})  '+text;r.append(t);math.append(r);om.append(math);p._p.append(om);self.md+=['```text',f'({self.eq}) '+text,'```','']
    def figure(self,path,caption):self.doc.add_picture(str(path),width=Inches(6.2));self.p(caption);self.md+=['!['+caption+'](../../final_evaluation/'+path.name+')','']
    def save(self):
        (ROOT/'reports/summaries/evaluation').mkdir(parents=True,exist_ok=True)
        self.doc.save(ROOT/'reports/summaries/evaluation/evaluation_tasks_report.docx');(ROOT/'reports/summaries/evaluation/evaluation_tasks_report.md').write_text('\n'.join(self.md),encoding='utf-8')

def run():
    selection=read(OUT/'selection_summary.json');lock=read(OUT/'lock.json');test=read(OUT/'test_metrics.json');boot=read(OUT/'bootstrap.json')
    audit=read(OUT/'manual_error_review.json');audit_selection=read(OUT/'audit_selection.json');fits=read(OUT/'fit_progress.json')
    baseline=read(ROOT/'reports/nucle_pos_baseline_v1/summary.json')
    baseline_validation={'Majority':baseline['majority_baseline_validation']['macro_f1'],'POS':baseline['pos_rule_validation']['macro_f1']}
    for e in lock['selected']:NAMES[e['id']]=e['model']+' '+e['variant']
    primary=['Majority','POS','LR_03','Tree_03'];allmodels=test['all']['models'];byid={e['id']:e for e in lock['selected']}
    overall=[];perclass=[]
    for ident,m in allmodels.items():
        p=m['per_label'];meanp=statistics.mean(v['precision'] for v in p.values());meanr=statistics.mean(v['recall'] for v in p.values())
        overall.append({'pipeline':ident,'name':NAMES[ident],'validation_macro_f1':byid[ident]['validation_macro_f1'] if ident in byid else baseline_validation[ident],'test_macro_f1':m['macro_f1'],'accuracy':m['accuracy'],'macro_precision':meanp,'macro_recall':meanr,'ci95_low':boot['slices']['all']['models'][ident]['ci95'][0],'ci95_high':boot['slices']['all']['models'][ident]['ci95'][1],'alias_of':byid.get(ident,{}).get('alias_of','')})
        for label,v in p.items():perclass.append({'pipeline':ident,'label':label,**v})
    delta=[]
    for e in lock['selected']:
        base=allmodels[e['model']+'_03']['per_label']
        for label in LABELS:delta.append({'pipeline':e['id'],'variant':e['variant'],'label':label,'delta_f1_vs_full':allmodels[e['id']]['per_label'][label]['f1']-base[label]['f1']})
    r=Report();r.h('Báo cáo thực hiện task: ablation bổ sung, test, bootstrap và audit lỗi',0)
    r.p('Ngày thực hiện task: 23/10/2026. Tài liệu ghi nhận công việc vừa chạy, đầu vào, code, cấu hình, kết quả và kiểm chứng. Đây là báo cáo task, không phải final report của dự án. Các nhận xét bên dưới giải thích bằng chứng thực nghiệm; việc tổng hợp lập luận và chốt câu trả lời RQ1–RQ3 trong báo cáo tổng kết được để lại cho bước sau theo yêu cầu người dùng.')
    r.h('1. Phạm vi task và đầu ra đã có')
    r.p('LR Full đạt test macro-F1 0,8163, CI95% [0,7876; 0,8382], vượt Majority 0,1184, POS 0,7106 và Tree Full 0,7996. Chênh LR–POS là 0,1057 [0,0829; 0,1288]; LR–Tree Full là 0,0166 [0,0018; 0,0301]. Đây là CI pointwise từ cùng 1.000 mẫu bootstrap theo 97 connected groups, không phải khoảng tin cậy từ nhiều seed huấn luyện.')
    r.p('Nhận xét từ bảng ablation: POS bổ sung mạnh cho mốc S+L; lexical vẫn hữu ích trên test. Bỏ toàn L làm LR giảm 0,0281, lớn hơn mức giảm khoảng 0,0103 trên validation. Morphology có tác dụng rõ hơn với Tree; CI của tác dụng M và D với LR chứa 0 nên chưa đủ để khẳng định cải thiện ổn định. Không diễn giải kết quả này là thông tin hình thái/dependency vô ích.')
    r.p('Nhận xét từ phân tích lỗi: hai nhóm nhầm lớn vẫn là Nn/ArtOrDet và Vform/SVA. LR cải thiện rõ ở các tập con cùng POS nhưng vẫn gặp span rộng, lỗi thiếu từ neo vào token khác từ loại, parse đáng ngờ và gold/edit không khớp định nghĩa nhãn đơn giản. Audit 40 ca là review có hỗ trợ AI, chưa phải adjudication của chuyên gia.')
    r.table(['Hạng mục','Trạng thái / bằng chứng'],[
        ['Ablation còn thiếu','40 fit grid mới Full−P/M, 12 refit cấu hình cũ; Full−D alias của S+L+P+M'],
        ['Replay trước khóa','12 checkpoint cũ tái tạo đúng mọi prediction validation; Full encoder khớp 13.420 train/validation rows; POS validation khớp'],
        ['Khóa và test','18 tên LR/Tree + Majority/POS, 16 checkpoint LR/Tree riêng; test 2.368 sample, 97 groups, 144 essays'],
        ['Bootstrap','1.000 lần, seed20261023; 30 contrast/slice đủ điều kiện, ngưỡng và missing-class được ghi'],
        ['Audit','25 lỗi LR + 5 lỗi Tree khi LR đúng + 10 mẫu LR đúng, giữ ID và ghi chú'],
        ['Tái lập','Fresh venv nạp lại checkpoint từ encoded test, predictions khớp; không phải chạy lại toàn grid trong fresh env'],
        ['BERT','Giữ kết quả validation/chẩn đoán cũ; không chạy thêm hoặc chấm BERT trên test']])
    r.h('2. Đầu vào và điều kiện thực hiện')
    r.p('Tái sử dụng preparation v2, grouped split seed20261013, feature artifact và sáu biến thể đã chạy. Không chạy lại preprocessing, đổi split, sửa label hoặc thêm BERT trong đợt task này. Câu gốc/gold span là đầu vào; correction chỉ dùng audit sau prediction. Chi tiết thiết kế dữ liệu/mô hình nằm ở các báo cáo trước.')
    r.table(['Input','Vai trò'],[['data_set/processed/nucle_preparation_v2','Câu/span và metadata nguồn'],['data_set/split/nucle_group_split','Train11052 / validation2368 / test2368; partition cố định'],['data_set/features/nucle_features_v1','Encoder train-only và Full feature matrix16180 chiều'],['reports/nucle_feature_ablation_v1','Configs/predictions validation sáu variants cũ để replay'],['docs/final_evaluation_protocol.md','Protocol ghi trước chạy; không sửa lại theo điểm test']])
    r.p('Feature NLP: spaCy3.8.7 + en_core_web_sm3.8.0. Numeric giữ giá trị gốc; giới hạn500 chỉ áp dụng pooled n-gram, không phải mọi lexical slot. D chứa thông tin POS/morph của head/subject, nên bỏ P/M chỉ bỏ nhóm feature tương ứng. Train1069 essays/1045 groups; validation142/99; test144/97.')
    r.h('3. Phép tính dùng trong các task')
    r.math('p(y=k|x)=exp(xᵀw_k+b_k)/Σ_j exp(xᵀw_j+b_j); ŷ=argmax_k p(y=k|x)')
    r.math('J(W,b)=[Σ_i α_(y_i)(−log p(y_i|x_i))/Σ_i α_(y_i)] + ||W||²_F/(2Cn)')
    r.p('LR dùng code hiện hành: α=1 hoặc balanced α_k=n/(5n_k); intercept không L2. PyTorch LBFGS strong_wolfe,10 threads,seed20261013,max_iter1000/max_eval1250,tolerance_grad1e−6/change1e−9. Change-tolerance không đồng nghĩa gradient-converged.')
    r.math('Gini(t)=1−Σ_k p²_(k,t); gain=Gini(t)−(w_L/w_t)Gini(L)−(w_R/w_t)Gini(R)')
    r.p('Tree giữ implementation Gini weighted khi balanced; nhãn lá là majority raw counts. Không đổi thuật toán sau test.')
    r.math('Precision_k=TP_k/(TP_k+FP_k); Recall_k=TP_k/(TP_k+FN_k)')
    r.math('F1_k=2TP_k/(2TP_k+FP_k+FN_k); MacroF1=(1/K)Σ_k F1_k')
    r.math('Accuracy=số prediction đúng/n')
    r.p('K=5 cho toàn tập/span bins, K=2 cho cặp gold; predictions vẫn năm lớp, ngoài cặp gây false negative. Zero denominator dùng0; subset rỗng ghi N/A. Công thức bootstrap ở mục8.')
    r.h('4. Task bổ sung ablation, tái tạo checkpoint và khóa')
    r.p('Sáu variants cũ chỉ lưu summary/predictions nên cần refit đúng config để có checkpoint. Tất cả 12 refit khớp validation predictions cũ ở10 threads. Full−D chính là S+L+P+M: giữ alias để trình bày RQ2, không chạy thêm20 fit và không tính hai bằng chứng độc lập. Hai variants mới chạy40 fit; tổng phiên52 fit, khoảng '+f"{selection['runtime_seconds']/60:.2f}"+' phút gồm kiểm/tạo encoder. Bốn LR C=10 mới chạm iteration limit và không eligible; mọi LR được chọn dừng change_tolerance.')
    r.p('Grid LR C={0,01;0,1;1;10} × class_weight={None,balanced}; Tree depth={4,8,12} × leaf={5,20} × class_weight. Chọn max macro-F1 validation trong eligible, hòa lấy cấu hình đầu theo grid. Từng variant tune riêng, nên đo pipeline theo cùng ngân sách tuning, không cô lập tác dụng feature ở cùng hyperparameters.')
    r.table(['Pipeline','Số chiều','Config chọn','Val F1'],[[NAMES[e['id']],e['feature_count'],json.dumps(e['configuration'],ensure_ascii=False),fmt(e['validation_macro_f1'])] for e in lock['selected']])
    r.p('Full_mask_span_lexical loại lexical identity của token giao span trước khi fit encoder masked riêng trên train; giữ context, POS/morph/dependency của câu gốc, không nối n-gram qua phần bị che. Nonlexical loại toàn L, không gọi structural-only.')
    r.p('Lock ghi lúc '+lock['locked_at']+'; test bắt đầu '+read(OUT/'test_scoring_started.json')['at']+'. SHA256 lock: '+sha(OUT/'lock.json')+'. Lock chứa hashes source/preparation/split/code/encoders/checkpoints/protocol và label order. Test scorer kiểm hash trước chạy. Không refit train+validation, không chọn model/seed/nhóm phân tích sau điểm test.')
    r.h('5. Task chấm test: kết quả và kiểm tra')
    r.table(['Model','Val F1','Test F1','CI95%','Accuracy','Macro P','Macro R'],[[x['name'],fmt(x['validation_macro_f1']),fmt(x['test_macro_f1']),ci([x['ci95_low'],x['ci95_high']]),fmt(x['accuracy']),fmt(x['macro_precision']),fmt(x['macro_recall'])] for x in overall if x['pipeline'] in primary])
    r.figure(OUT/'benchmark_ci.png','Hình 1. Macro-F1 test và CI95% theo connected group; model/config đã khóa trước test.')
    r.table(['So sánh','Chênh F1','CI95%'],[[d['contrast'],fmt(d['difference']),ci(d['ci95'])] for d in boot['differences'][:4]])
    r.p('Diễn giải bảng test (bằng chứng cho RQ1): cả LR và Tree vượt Majority rõ rệt trong task known-span này. LR Full vượt POS và Tree Full; các CI paired chính không chứa0. So LR/Tree dùng cùng Full, không chọn Tree variant theo điểm test. Tree masking có điểm test nhỉnh Full rất nhỏ, nhưng không dùng quan sát đó để thay model đã chọn.')
    r.p('Macro-F1 LR tăng so validation trong khi accuracy giảm: metric macro cân bằng năm lớp, còn accuracy chịu trọng số theo support. F1 Nn test0,7781 thấp hơn validation khoảng0,8383, trong khi Vform/SVA tốt hơn; đây không phải mâu thuẫn metric. Không suy test cao hơn là hết overfitting.')
    r.table(['Nhãn','Support','LR P','LR R','LR F1','POS F1','Tree Full F1'],[[l,allmodels['LR_03']['per_label'][l]['support'],fmt(allmodels['LR_03']['per_label'][l]['precision']),fmt(allmodels['LR_03']['per_label'][l]['recall']),fmt(allmodels['LR_03']['per_label'][l]['f1']),fmt(allmodels['POS']['per_label'][l]['f1']),fmt(allmodels['Tree_03']['per_label'][l]['f1'])] for l in LABELS])
    r.h('6. Task đối chiếu ablation/masking')
    r.table(['Variant','LR val','LR test','Tree val','Tree test'],[[e['variant'],fmt(e['validation_macro_f1']),fmt(allmodels[e['id']]['macro_f1']),fmt(byid[e['id'].replace('LR','Tree')]['validation_macro_f1']),fmt(allmodels[e['id'].replace('LR','Tree')]['macro_f1'])] for e in lock['selected'] if e['model']=='LR'])
    r.table(['Đối chiếu','Chênh test F1','CI95%'],[[d['contrast'],fmt(d['difference']),ci(d['ci95'])] for d in boot['differences'][4:20]])
    r.p('Thêm POS vào S+L tăng LR0,0766 và Tree0,1670; CI đều dương. Trên nền Full, bỏ nhóm P chỉ giảm khoảng0,018 vì feature groups chồng thông tin. Hai phép đo trả lời hai câu hỏi có điều kiện khác nhau, không mâu thuẫn.')
    r.p('Với LR, che lexical span giảm0,0134 [0,0049;0,0201], bỏ toàn L giảm0,0281 [0,0142;0,0394]. Lexical giúp ích trên test dù POS/nonlexical đã mạnh. Với Tree, chênh Full−masked là−0,0010 với CI chứa0; không khẳng định masking làm Tree tốt hơn. Không gọi masking là kiểm chứng model hiểu ngữ pháp.')
    r.p('Thêm M sau P giúp LR0,0134 nhưng CI[−0,0004;0,0256] chứa0; bỏ M khỏi Full giảm0,0040 với CI cũng chứa0. Tree tăng0,0447 khi thêm M, giảm0,0171 khi bỏ M khỏi Full, các CI dương. Cấu trúc model và tuning tạo khác biệt; không gán hiệu ứng chung giống nhau cho LR/Tree.')
    r.p('Thêm D giúp LR0,0060 và Tree0,0029, cả CI chứa0. Chưa đủ bằng chứng lợi ích tăng thêm ổn định của D trong cấu hình hiện tại; không kết luận chủ ngữ/ngữ cảnh toàn câu vô ích. Extractor chỉ lấy subject child trực tiếp và parser có thể sai trên câu lỗi.')
    for family in ['LR','Tree']:
        r.h('F1 từng lớp và Δ so Full — '+family,2)
        r.table(['Variant',*LABELS],[[e['variant'],*[fmt(allmodels[e['id']]['per_label'][l]['f1'])+' ('+f"{allmodels[e['id']]['per_label'][l]['f1']-allmodels[family+'_03']['per_label'][l]['f1']:+.4f}"+')' for l in LABELS]] for e in lock['selected'] if e['model']==family])
    r.p('Diễn giải ablation (bằng chứng cho RQ2): P và L có bằng chứng hỗ trợ rõ ở LR; M thể hiện tác dụng rõ hơn với Tree; tác dụng tăng thêm của D và của M trong LR chưa chắc chắn. Các Δ trong bảng từng lớp là descriptive, không có CI từng lớp đã hiệu chỉnh. Full−D lặp S+L+P+M chỉ để thuận đối chiếu, không nhân đôi bằng chứng.')
    r.h('7. Task phân tích span, cặp nhãn, POS và audit lỗi')
    r.table(['Subset','n','Groups','POS F1','LR F1','Tree F1','Δ LR−POS CI95%'],[[name,v['support'],v['groups'],fmt(v['models']['POS']['macro_f1']),fmt(v['models']['LR_03']['macro_f1']),fmt(v['models']['Tree_03']['macro_f1']),ci(next((d['ci95'] for d in boot['differences'] if d['slice']==name and d['a']=='LR_03' and d['b']=='POS'),None))] for name,v in test.items() if name!='all'])
    r.p('Span một token n2102, nhiều token n266; LR macro5 tương ứng0,8320 và0,5370. Δ LR−POS trên multi-token có CI chứa0, nên không khẳng định cải thiện chắc chắn ở mọi span rộng. Span rỗng support0 ghi N/A; không loại mẫu chỉ vì thiếu POS.')
    r.p('Gold Vform/SVA n445: LR macro2≈0,8021 so POS≈0,6095; trên primary POS VERB n220, LR≈0,8416 so POS≈0,3548. Nn/ArtOrDet n1561: LR≈0,8453 so POS≈0,7663; giao NOUN n682: LR≈0,8142 so POS≈0,4044. Những kết quả này cho thấy đặc trưng ngoài signature POS có ích, chưa chỉ ra riêng feature nào là nguyên nhân.')
    r.p('Primary POS chỉ có giá trị khi span đúng một token theo contract. NOUN slice không phải parser-confirmed noun phrase. Cặp Nn/ArtOrDet@DET có macro2≈0,4827 dù phần lớn là ArtOrDet: trung bình vẫn tính cả Nn và giữ prediction ngoài cặp; không đọc điểm này như accuracy thấp tương ứng. Subset ít mẫu không có CI vẫn được giữ trong bảng.')
    r.figure(OUT/'confusion_matrices.png','Hình 2. Confusion test của LR Full và POS; hàng gold, cột prediction.')
    c=allmodels['LR_03']['confusion'];pairs=sorted([(c[i][j],LABELS[i],LABELS[j]) for i in range(5) for j in range(5) if i!=j],reverse=True)
    r.table(['Gold → prediction','Số lỗi'],[[a+' → '+b,n] for n,a,b in pairs[:10]])
    r.p('Nn→ArtOrDet134 và ArtOrDet→Nn59 là hai hướng nhầm lớn nhất; Vform→SVA29, SVA→Vform18. Số lượng chịu ảnh hưởng support nên đối chiếu per-class recall/F1, không chỉ xếp count. Chưa có span rỗng nên không thể nói phần lớn Nn→ArtOrDet nằm trong nhóm span rỗng; audit cho thấy missing-word edits có thể có span không rỗng.')
    r.h('7.1 Audit sau prediction và cách giải thích',2)
    r.p('Chọn trước khi đọc correction:25 lỗi LR round-robin gold/pred,5 lỗi Tree khi LR đúng,10 controls LR đúng. Seed20261023; không chọn ngẫu nhiên đại diện toàn corpus. audit_selection.json lưu IDs; private_error_audit_packet.json giữ câu/correction/parse local; manual_error_review.json chỉ giữ ID/nhận xét không trích câu. Tất cả40 ca được đọc sau prediction. Confidence là mức chắc chắn nhận xét, không phải xác suất model.')
    r.table(['Quan sát','ID minh họa','Ý nghĩa và giới hạn'],[
        ['Thiếu từ, span không rỗng','38087,24747,17826,348,44388,34548','POS của token neo có thể khác POS của từ thiếu; không phải lỗi offset'],
        ['Parser/POS đáng ngờ','24526,9351,41235,35455,211,37624','Noun/verb/adjective bị gán khác cách đọc; cần chuyên gia để xác nhận lỗi parser'],
        ['Partial-token','17814','Ký tự gold thuộc token ADP bao phủ, không phản ánh đúng loại edit số danh từ'],
        ['Span rộng/nhiều edit','19312,34483,32904','Một nhãn bao phủ đổi determiner/number hoặc cấu trúc mệnh đề'],
        ['Phối hợp/subject gián tiếp','18531,36821,17311','D chỉ tìm subject child của target, không truyền qua auxiliary/conj'],
        ['Gold/edit cần adjudication','34539,12324,26635,32705,37909,18501,201','Không tự sửa nhãn. Prediction khớp gold cũng chưa chứng minh phán đoán ngữ pháp đúng'],
        ['LR vượt POS ở missing determiner','11916,44353,6079','Có ca LR phân biệt được dù token NOUN; chưa xác định nhân quả của feature']])
    r.p('Diễn giải audit (bằng chứng cho RQ3): nhầm lẫn liên quan taxonomy trong cụm danh từ/động từ, độ rộng span, token neo của missing-word edits, chất lượng parse và sự phù hợp của gold với edit. Không suy tỷ lệ từng nguyên nhân từ audit chủ đích, không loại các ca này rồi báo lại test đẹp hơn.')
    r.h('8. Task bootstrap: công thức và kiểm soát')
    r.math('m^(b) ~ Multinomial(G; 1/G,…,1/G), G=97; C_model^(b)=Σ_g m_g^(b) C_(g,model)')
    r.math('Δ^(b)=MacroF1(C_A^(b))−MacroF1(C_B^(b)); CI95%=[Q_0.025(Δ),Q_0.975(Δ)]')
    r.p('C_(g,model) là confusion5×5 của connected group; cùng m^(b) cho mọi model/subset. Group có nhiều essay vẫn là một đơn vị bootstrap. Với subset, lọc annotation bên trong group sau khi lấy cùng group weights. Không resample riêng từng annotation; điều đó sẽ bỏ phụ thuộc trong bài.')
    r.p('Ngưỡng subset CI: ≥20 samples, ≥10 groups và có đủ labels cần tính. Replicate thiếu lớp vẫn tính trên fixed labels (zero denominator→0), ghi missing-class count; replicate subset empty bị bỏ và CI chỉ báo nếu≥950/1000 hợp lệ. Các CI pointwise không hiệu chỉnh multiplicity; không có p-value hoặc kết luận significance đồng thời.')
    r.table(['Subset','Eligible CI','Valid','Empty','Missing-class'],[[name,str(v['eligible']),v['models'].get('LR_03',{}).get('valid_replicates','N/A'),v['models'].get('LR_03',{}).get('empty_replicates','N/A'),v['models'].get('LR_03',{}).get('missing_class_replicates','N/A')] for name,v in boot['slices'].items()])
    r.h('9. Phạm vi BERT trong đợt task này')
    r.p('Không chạy thêm BERT và không chấm BERT trên test. Các kết quả BERT-small validation/học thuộc/gradient/span thuộc đợt trước, xem contextual_model_validation_report.docx và bert_learning_diagnostics_report.docx. Chúng được giữ như thử nghiệm chưa thành công, không dùng kết luận ngữ cảnh vô ích.')
    r.h('10. Task kiểm chứng, tái lập và bàn giao artifact')
    r.table(['File','Input → output / vai trò'],[
        ['src/training/nucle_final_selection.py','Feature/split/old predictions → train-only encoder,52fit,16checkpoints,validation replay,lock'],
        ['src/evaluation/nucle_final_evaluation.py','Lock + test → saved predictions/metrics; bootstrap từ predictions; checkpoint replay'],
        ['src/evaluation/nucle_final_error_audit.py','Saved predictions → deterministic IDs → private source/correction/parse packet'],
        ['src/reporting/build_evaluation_tasks_report.py','JSON/CSV + manual notes → báo cáo task DOCX/Markdown; tái dùng hình đã có'],
        ['docs/final_evaluation_protocol.md','Protocol ghi trước test, quy tắc alias/selection/subsets/bootstrap'],
        ['reports/final_evaluation/lock.json','Hash input/code/encoders/checkpoints, thời điểm và cấu hình'],
        ['test_predictions.csv, test_metrics.json','Prediction từng sample/pipeline và metrics theo subsets'],
        ['bootstrap.json, bootstrap_group_weights.npy','CI paired, số lượt thiếu lớp và resampling weights'],
        ['overall_results.csv, per_class_results.csv, per_class_deltas.csv','Bảng tổng thể, từng lớp và Δ so Full'],
        ['manual_error_review.json','40 ghi chú không có câu nguồn; private packet giữ riêng local'],
        ['reproducibility/README.md','Lệnh chạy, dependencies và phạm vi kiểm tái lập']])
    r.p('13 tests phù hợp đã qua, gồm5 tests mới cho metric ngoài cặp, group bootstrap, serialization encoder, loại nhóm/alias và từ chối lock bị sửa. Các tests model/feature hiện có cũng qua. Checkpoint reload18 tên LR/Tree không bất đồng prediction. Fresh venv dùng Python3.13.7, NumPy2.5.3, PyTorch2.14.0+cpu; inference trên encoded test khớp tuyệt đối. Kiểm này không thay toàn bộ re-training, spaCy reparse hoặc cài lại mọi dependency trong môi trường sạch.')
    r.p('Gói hỗ trợ tái lập task đóng code/protocol/summary/hash/hướng dẫn; không chứa corpus, câu/correction audit, encoder vocab hoặc checkpoint. Người tái lập đầy đủ cần quyền truy cập NUCLE và artifact nguồn phù hợp. Không upload/push trong lượt này. Restore gói vào thư mục mới được kiểm hash; manifest cũ giữ nguyên, report chỉ dẫn đến artifact mới.')
    r.h('11. Hạn chế và công việc để lại')
    r.table(['Hạng mục','Trạng thái'],[['Các task thí nghiệm','Ablation bổ sung, khóa/test, bootstrap và audit đã chạy; số liệu giữ nguyên'],['Báo cáo task này','Ghi mục tiêu, input/code/output, kết quả, giải thích và kiểm chứng của đợt vừa làm'],['Final report của dự án','Chưa thực hiện theo yêu cầu mới; chưa tổng hợp thành báo cáo nộp cuối hoặc chốt lập luận RQ'],['Slide tổng kết','Chưa chốt; bản đã tạo trước đây được lưu như bản nháp lịch sử, không là đầu ra của task report'],['Tái lập end-to-end','Fresh venv mới kiểm checkpoint inference từ encoded test và metric; chưa chạy toàn preprocessing/grid từ raw'],['Audit chuyên gia','40ca do AI hỗ trợ review, chưa adjudication độc lập']])
    r.p('Một grouped split, một seed mô hình cổ điển; CI pointwise chưa hiệu chỉnh nhiều so sánh. Tập audit chủ đích không đại diện tỷ lệ nguyên nhân toàn corpus. Token/parser tự động và gold/edit mơ hồ vẫn là giới hạn. Không sửa model/labels theo test đã xem; các hướng mở rộng cần đánh giá mới.')
    r.h('12. Tài liệu và artifact đối chiếu')
    r.p('Kế hoạch: note.md. Protocol trước chạy: docs/final_evaluation_protocol.md. Phương pháp dữ liệu/mô hình đã có: docs/note_project_guide.md, reports/summaries/data/data_preprocessing_report.docx, reports/summaries/classical/modeling_ablation_results_report.docx. Nguồn học thuật: reference.md. Task report không thay thế các tài liệu này hay viết lại toàn bộ related work. Bảng số liệu được lấy từ reports/final_evaluation; không dùng kết quả chưa chạy.')
    r.h('Phụ lục A. Ghi chú audit theo ID')
    selected={str(x['annotation_id']):x for x in audit_selection['rows']}
    for n in audit['notes']:
        a=selected[n['annotation_id']];r.h('ID '+n['annotation_id']+' — '+a['gold']+'; LR='+a['LR']+', Tree='+a['Tree']+', POS='+a['POS'],2)
        r.p('Nhóm: '+n['category']+'; confidence: '+n['confidence']+'. '+n['explanation'])
    r.h('Phụ lục B. Danh mục toàn bộ fit và CI')
    r.p('52 dòng fit (gồm cả cấu hình không eligible) ở fit_progress.json. 20 pipeline rows ở overall_results.csv;100 dòng metric từng lớp ở per_class_results.csv;90 Δ class/variant ở per_class_deltas.csv. Alias được gắn rõ; không bỏ kết quả âm. Toàn bộ CI và số lượt thiếu lớp ở bootstrap.json. Các artifact này là nguồn dựng lại bảng, không phải nhập số tay.')
    r.save()
    print('TASK REPORT',ROOT/'reports/summaries/evaluation/evaluation_tasks_report.docx','equations',r.eq)
if __name__=='__main__':run()
