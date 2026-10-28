"""Small Markdown-to-DOCX/PDF exporter for the project's Vietnamese reports."""
from __future__ import annotations

from html import escape
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, KeepTogether, LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle


def fonts():
    from matplotlib import get_data_path
    folder = Path(get_data_path()) / "fonts/ttf"
    for name, file in (("DejaVu", "DejaVuSans.ttf"), ("DejaVu-Bold", "DejaVuSans-Bold.ttf"),
                       ("DejaVu-Oblique", "DejaVuSans-Oblique.ttf"), ("DejaVu-BoldOblique", "DejaVuSans-BoldOblique.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(folder / file)))
    pdfmetrics.registerFontFamily("DejaVu", normal="DejaVu", bold="DejaVu-Bold", italic="DejaVu-Oblique", boldItalic="DejaVu-BoldOblique")


def plain(text):
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return text.replace("**", "").replace("`", "").replace("\\|", "|")


def inline(text):
    text = escape(text.replace("\\|", "|"))
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r'<link href="\2" color="#087f8c">\1</link>', text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", text)
    return text.replace("`", "")


def blocks(path):
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i]); i += 1
            yield "code", "\n".join(code)
        elif line.startswith("|"):
            table = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", lines[i].strip().strip("|"))]
                if not all(re.fullmatch(r":?-+:?", c.replace(" ", "")) for c in cells):
                    table.append(cells)
                i += 1
            if len({len(row) for row in table}) != 1:
                raise ValueError("Inconsistent Markdown table columns")
            yield "table", table
            continue
        elif line.startswith("#"):
            m = re.match(r"^(#+)\s+(.*)", line)
            yield "heading", (len(m[1]), m[2])
        elif line.startswith("!["):
            m = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
            if not m:
                raise ValueError(f"Unsupported image syntax: {line}")
            image = (path.parent / m[2]).resolve()
            if not image.exists():
                raise FileNotFoundError(image)
            yield "image", (m[1], image)
        else:
            yield "paragraph", line
        i += 1


def export_docx(markdown, output):
    document = Document()
    section = document.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2), Cm(2)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    for level, size in ((1, 16), (2, 13), (3, 11)):
        style = document.styles[f"Heading {level}"]
        style.font.name, style.font.size, style.font.color.rgb = "Arial", Pt(size), RGBColor.from_string("153447")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("NUCLE • 28/10/2026 • ")
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for kind, value in blocks(markdown):
        if kind == "heading":
            level, text = value
            document.add_heading(plain(text), level=min(level, 3))
        elif kind == "table":
            table = document.add_table(rows=1, cols=len(value[0]))
            table.style = "Light Shading Accent 1"
            for cell, text in zip(table.rows[0].cells, value[0], strict=True):
                cell.text = plain(text)
            header = OxmlElement("w:tblHeader"); table.rows[0]._tr.get_or_add_trPr().append(header)
            for row in value[1:]:
                for cell, text in zip(table.add_row().cells, row, strict=True):
                    cell.text = plain(text)
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(2)
                        for run in p.runs:
                            run.font.size = Pt(8 if len(value[0]) > 5 else 9)
            document.add_paragraph()
        elif kind == "image":
            caption, image = value
            p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(image), width=Inches(6.45))
            caption_p = document.add_paragraph(plain(caption)); caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in caption_p.runs: run.font.size = Pt(9); run.italic = True
        elif kind == "code":
            p = document.add_paragraph(value)
            for run in p.runs: run.font.name = "Consolas"; run.font.size = Pt(8)
        else:
            text = plain(value)
            p = document.add_paragraph(text)
            if text.startswith("Nguồn:"):
                for run in p.runs: run.font.size = Pt(8); run.italic = True
    document.core_properties.title = "Phân loại năm loại lỗi ngữ pháp tiếng Anh trên NUCLE"
    document.core_properties.subject = "Kết quả cổ điển và BERT-small đã khóa, cập nhật 28/10/2026"
    document.core_properties.author = "Nhóm nghiên cứu — tên và đóng góp cần xác nhận"
    document.save(output)


def export_pdf(markdown, output):
    fonts()
    width = A4[0] - 90
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="BodyVN", fontName="DejaVu", fontSize=9, leading=13.3, spaceAfter=6, alignment=TA_JUSTIFY))
    styles.add(ParagraphStyle(name="SmallVN", fontName="DejaVu", fontSize=7.4, leading=10, spaceAfter=5))
    styles.add(ParagraphStyle(name="CellVN", fontName="DejaVu", fontSize=7.3, leading=10, splitLongWords=True))
    styles.add(ParagraphStyle(name="CaptionVN", fontName="DejaVu", fontSize=8, leading=11, alignment=TA_CENTER, spaceAfter=7))
    for level, size in ((1, 16), (2, 12), (3, 10.5)):
        styles.add(ParagraphStyle(name=f"HVN{level}", fontName="DejaVu-Bold", fontSize=size, leading=size * 1.3,
                                 textColor=colors.HexColor("#153447"), spaceBefore=12, spaceAfter=7, keepWithNext=True))
    story = []
    for kind, value in blocks(markdown):
        if kind == "heading":
            level, text = value
            story.append(Paragraph(inline(text), styles[f"HVN{min(level, 3)}"]))
        elif kind == "table":
            n = len(value[0])
            ratios = [1] * n
            if n == 4 and any(len(c) > 90 for row in value for c in row):
                ratios = [1, 1.5, 1.5, 2]
            widths = [width * r / sum(ratios) for r in ratios]
            cells = [[Paragraph(inline(text), styles["CellVN"]) for text in row] for row in value]
            table = LongTable(cells, colWidths=widths, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e0edf1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f8fa")]),
                ("GRID", (0, 0), (-1, -1), .3, colors.HexColor("#cbd5e1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.extend([table, Spacer(1, 8)])
        elif kind == "image":
            caption, path = value
            image = Image(str(path))
            scale = min(width / image.imageWidth, 270 / image.imageHeight)
            image.drawWidth, image.drawHeight = image.imageWidth * scale, image.imageHeight * scale
            story.append(KeepTogether([image, Spacer(1, 4), Paragraph(inline(caption), styles["CaptionVN"])]))
        elif kind == "code":
            for line in value.splitlines():
                story.append(Paragraph(escape(line), styles["SmallVN"]))
        else:
            story.append(Paragraph(inline(value), styles["SmallVN"] if value.startswith("Nguồn:") else styles["BodyVN"]))

    def footer(canvas, doc):
        canvas.setFont("DejaVu", 7)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawString(45, 23, "NUCLE • Kết quả đã khóa • Cập nhật 28/10/2026")
        canvas.drawRightString(A4[0] - 45, 23, str(doc.page))

    SimpleDocTemplate(str(output), pagesize=A4, leftMargin=45, rightMargin=45, topMargin=40, bottomMargin=40,
                      title="Phân loại lỗi ngữ pháp NUCLE — báo cáo tổng kết", author="Nhóm nghiên cứu — chờ xác nhận thành viên").build(
        story, onFirstPage=footer, onLaterPages=footer)
