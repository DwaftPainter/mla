"""Generate editable PPTX, matching-layout PDF and Vietnamese speaker notes."""
from __future__ import annotations

import math
from pathlib import Path
from textwrap import wrap

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.lib.utils import ImageReader

from src.evaluation.bert_final_analysis import ROOT, csv_read, dump, read
from src.reporting.document_export import fonts

FINAL = ROOT / "reports/summaries/final"
ANALYSIS = ROOT / "reports/bert_final_analysis"
NAVY, TEAL, ORANGE, INK, GRAY, LIGHT = "153447", "087F8C", "C45B39", "183447", "536778", "F3F7F9"
WIDTH, HEIGHT = 13.333, 7.5


def specs():
    overall = csv_read(ANALYSIS / "overall_results.csv")
    return [
        {"title": "Phân loại lỗi ngữ pháp trên NUCLE", "subtitle": "Biết vị trí lỗi, mô hình có nhận ra đúng loại lỗi không?", "bullets": [
            "15.788 mẫu thuộc năm nhãn; chia tập theo nhóm bài viết.",
            "Macro-F1 test: LR Full 0,8163; BERT ở seed chính 0,8146.",
            "Nhóm thực hiện: [điền họ tên và MSSV]"],
         "notes": "Bài trình bày tập trung vào một câu hỏi cụ thể: khi đã biết đoạn nào trong câu có lỗi, mô hình có thể phân loại lỗi đó tốt đến đâu? Đây chưa phải hệ tự tìm và sửa lỗi. Bộ slide dự kiến khoảng 10 phút. Nhóm cần bổ sung họ tên và MSSV trước khi sử dụng để nộp bài.", "seconds": 35},
        {"title": "Bài toán và ba câu hỏi nghiên cứu", "bullets": [
            "Đầu vào là câu gốc và vị trí bắt đầu, kết thúc của đoạn lỗi.",
            "Đầu ra: ArtOrDet, Prep, Nn, Vform hoặc SVA.",
            "RQ1: LR/Tree so với Majority và POS rule tốt đến đâu?",
            "RQ2: từ vựng, POS, hình thái và quan hệ cú pháp có ích ra sao?",
            "RQ3: mô hình còn nhầm những loại lỗi nào và trong trường hợp nào?"],
         "notes": "Ba câu hỏi đi từ hiệu quả tổng thể đến vai trò của thông tin đầu vào, rồi đến các lỗi còn lại. Nhãn dùng để học, còn bản sửa chỉ được đọc khi rà soát sau dự đoán; chúng không nằm trong đầu vào dự đoán. BERT được bổ sung sau để kiểm tra một cách biểu diễn khác trên cùng bài toán.", "seconds": 45},
        {"title": "Dữ liệu và tránh rò rỉ", "bullets": [
            "Từ 44.912 annotation trong 1.397 bài, giữ 15.788 mẫu thuộc năm nhãn.",
            "Train 11.052 • Validation 2.368 • Test 2.368.",
            "Các bài có nội dung trùng được gom cùng nhóm trước khi chia tập.",
            "Encoder chỉ học từ train; test có 144 bài thuộc 97 nhóm.",
            "ArtOrDet chiếm khoảng 42%, nên cần xem macro-F1 bên cạnh accuracy."],
         "notes": "Một câu có thể tạo nhiều mẫu nếu chứa nhiều annotation. Nếu chia ngẫu nhiên từng dòng, các mẫu có liên hệ có thể xuất hiện ở cả train và test. Chia theo nhóm giúp hạn chế tình huống này. Cũng vì sự phụ thuộc đó, khoảng tin cậy được lấy mẫu lại theo nhóm bài viết.", "seconds": 50},
        {"title": "Từ dữ liệu đến kết quả: ba phần công việc", "bullets": [
            "1. Dữ liệu và đặc trưng: đọc SGML, kiểm offset, chia tập, chạy spaCy và mã hóa S/L/P/M/D.",
            "2. Mô hình: huấn luyện các đối chứng, LR, Tree và BERT-small.",
            "3. Đánh giá: khóa lựa chọn, chấm test, tính CI, đọc lỗi và viết báo cáo.",
            "BERT nhận câu có đánh dấu span; LR/Tree nhận vector đặc trưng."],
         "notes": "Đây là cách chia nội dung thành ba phần để giải thích quy trình. Phần tạo đặc trưng thuộc bước dữ liệu. Biểu diễn Full của LR và Tree có 16.180 chiều; BERT học biểu diễn từ câu có marker, không nối thêm vector S/L/P/M/D. Phân công theo tên cần được xác nhận từ công việc thực tế.", "seconds": 45},
        {"title": "Thiết kế mô hình và lựa chọn", "bullets": [
            "LR phân loại đa lớp; Tree chia nhánh theo Gini với giới hạn độ sâu và kích thước lá.",
            "Chọn cấu hình bằng validation; mô hình cuối vẫn chỉ được học từ train.",
            "Bỏ hoặc che từng nhóm đặc trưng để xem điểm thay đổi thế nào.",
            "BERT-small có 2 layer, hidden size 128 và 2 head; fine-tune toàn bộ tham số."],
         "notes": "Các biến thể đặc trưng được chọn tham số riêng, nên chênh lệch chưa phải hiệu ứng nhân quả của một nhóm thông tin ở cùng bộ tham số. POS rule có bảng tra học từ train. LR cũng có dependency từ cả câu, vì thế đối chiếu với BERT không đơn giản là mô hình có ngữ cảnh và mô hình không có ngữ cảnh.", "seconds": 45},
        {"title": "Từ BERT ban đầu đến đợt mở rộng", "bullets": [
            "Đợt đầu: pooled CLS, 5 epochs → validation F1 0,1984 ± 0,0030.",
            "Kiểm trên 100 mẫu train cho thấy mô hình có thể học thuộc và nhận được gradient.",
            "Đợt mới: 2 cách lấy biểu diễn × 2 head LR × 3 seed = 12 lượt; encoder LR 5e−5.",
            "Theo dõi train và validation qua tối đa 20 epochs, dừng sớm với patience 4.",
            "Hai đợt khác cả cách lấy biểu diễn lẫn thiết lập huấn luyện."],
         "notes": "Thử 100 mẫu giúp kiểm tra khả năng tối ưu hóa, không cho biết mô hình sẽ làm tốt trên dữ liệu chưa học. Vì vậy, đợt mở rộng bắt đầu lại từ pretrained và đánh giá trên toàn validation. Khi giải thích cải thiện, cần nhớ ngoài readout còn có thời gian huấn luyện và learning rate của head thay đổi.", "seconds": 50},
        {"title": "Chọn BERT bằng validation, không chọn seed test", "table": [
            ["Readout", "Head LR", "Val macro-F1 mean ± std"],
            ["CLS", "1e−4", "0,7164 ± 0,0061"], ["CLS", "1e−3", "0,7227 ± 0,0147"],
            ["Span mean", "1e−4", "0,8183 ± 0,0011"], ["Span mean", "1e−3", "0,8177 ± 0,0022"]],
         "bullets": ["Chọn span mean và head LR 1e−4. Seed chính là 20261013, lấy checkpoint epoch 18.",
                     "Test trung bình ba seed: 0,8182 ± 0,0038; độ lệch chuẩn này khác CI bootstrap."],
         "notes": "Cấu hình được chọn bằng kết quả trung bình của cùng ba seed trên validation. Seed chính đã cố định từ trước, nên không thay bằng seed 20261015 dù seed đó có điểm test cao nhất. Span mean lấy trung bình biểu diễn của các token trong đoạn lỗi sau khi encoder đã xử lý cùng ngữ cảnh.", "seconds": 50},
        {"title": "Kết quả test trên cùng split", "image": ANALYSIS / "figures/benchmark_ci.png", "bullets": [
            "LR Full 0,8163 • BERT primary 0,8146 • Tree Full 0,7996 • POS 0,7106.",
            "BERT có accuracy 0,8438, cao hơn LR 0,8357, nhưng macro-F1 hơi thấp hơn."],
         "notes": "Các điểm trên hình đều dùng cùng 2.368 mẫu test. Điểm BERT là của seed chính đã khóa, còn trung bình ba seed được báo riêng. Accuracy đếm mọi mẫu ngang nhau, trong khi macro-F1 cho mỗi lớp trọng số ngang nhau; vì thế hai chỉ số có thể dẫn đến thứ tự khác nhau.", "seconds": 50},
        {"title": "Chênh lệch có chắc chắn không?", "image": ANALYSIS / "figures/paired_differences_focus.png", "bullets": [
            "BERT−LR: −0,0017; CI95% [−0,0184; 0,0182] → chưa đủ bằng chứng hơn/kém.",
            "CI BERT−Tree cũng chứa 0; CI BERT−POS nằm trên 0. Dùng 1.000 lượt lấy lại 97 nhóm."],
         "notes": "Điểm cần nhìn là CI của chênh lệch, không phải hai CI riêng có chồng nhau hay không. Cùng một lượt lấy lại nhóm được dùng cho mọi mô hình. CI chứa 0 nghĩa là chưa đủ bằng chứng xếp hạng chắc chắn, không chứng minh hai mô hình bằng nhau. Phép tính cũng chưa bao quát biến thiên do split mới hoặc nhiều seed huấn luyện.", "seconds": 55},
        {"title": "Điểm mạnh và điểm yếu theo lớp", "image": ANALYSIS / "figures/per_class_f1.png", "bullets": [
            "BERT cao hơn LR ở bốn lớp. Với Nn, F1 là 0,8098 so với 0,7781.",
            "SVA là điểm yếu của BERT: 0,7392 so với LR 0,8106."],
         "notes": "Biểu đồ theo lớp cho thấy hai mô hình có những đánh đổi khác nhau dù điểm trung bình gần nhau. SVA đáng được đọc kỹ vì chênh lệch lớn hơn các lớp còn lại. Tuy nhiên, chỉ nhìn dự đoán chưa đủ để kết luận attention hay parser là nguyên nhân gây sai.", "seconds": 45},
        {"title": "Audit lỗi và giới hạn suy luận", "bullets": [
            "Cùng đúng 1.860 ca; BERT đúng riêng 138, LR đúng riêng 119; cùng sai 251.",
            "Đọc kỹ 40 ca: 12 theo mỗi chiều sai khác, 10 cùng sai và 6 ca cùng đúng.",
            "Lỗi thiếu từ vẫn có thể neo vào đoạn không rỗng; một số span chỉ là mảnh từ.",
            "Nhãn và bản sửa ở một số ca cần chuyên gia xem lại; gold vẫn giữ nguyên.",
            "Ghi chú AI hỗ trợ đọc lỗi, chưa thay đánh giá độc lập hoặc cho biết tỷ lệ toàn corpus."],
         "notes": "Có thể dùng CSV trên máy để giải thích một ca đã tự đọc. ID 38155 là mảnh từ nhưng cả hai mô hình vẫn đúng, giúp tránh suy rằng mọi partial-token đều gây lỗi. ID 34539 cho thấy cần xem quan hệ giữa nhãn và bản sửa. Các ca được chọn có chủ đích, không phải mẫu ngẫu nhiên đại diện corpus.", "seconds": 50},
        {"title": "Kết luận và bàn giao", "bullets": [
            "LR với đặc trưng ngôn ngữ là đối chứng mạnh; BERT span mean đạt điểm gần LR.",
            "CI chưa cho thấy BERT vượt LR; cả kết quả đầu tiên và đợt mở rộng đều được giữ lại.",
            "Cần dữ liệu đánh giá mới để kiểm tra khả năng khái quát hóa độc lập.",
            "Báo cáo có mã, bảng, hình, kiểm chứng và CSV đọc lỗi để đối chiếu.",
            "Trước khi nộp: hoàn thiện tên nhóm, đóng góp thực tế và khai báo AI."],
         "notes": "Kết quả cho thấy cách biểu diễn và thiết lập huấn luyện đều đáng quan tâm, nhưng chưa cho phép tuyên bố mô hình hiểu ngữ pháp đầy đủ hoặc tốt trên mọi corpus. Đợt mở rộng dùng lại test đã biết từ các mô hình cổ điển, nên nghiên cứu tiếp nên có dữ liệu bên ngoài. Phần còn lại của nhóm là xác nhận thông tin và tự kiểm tra tài liệu trước khi nộp.", "seconds": 45},
    ]


def textbox(slide, x, y, w, h, text, size=22, color=INK, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.text_frame.word_wrap = True
    shape.text_frame.margin_left = shape.text_frame.margin_right = 0
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = text
    for run in paragraph.runs:
        run.font.name = "Arial"; run.font.size = Pt(size); run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
    return shape


def wrap_pdf(text, size, width):
    words = text.split()
    lines, current = [], ""
    for word in words:
        candidate = (current + " " + word).strip()
        if pdfmetrics.stringWidth(candidate, "DejaVu", size) > width and current:
            lines.append(current); current = word
        else:
            current = candidate
    if current: lines.append(current)
    return lines


def pdf_text(pdf, x, y, text, size=20, color=INK, bold=False, width=800, leading=None):
    pdf.setFillColor(colors.HexColor("#" + color))
    pdf.setFont("DejaVu-Bold" if bold else "DejaVu", size)
    lines = wrap_pdf(text, size, width)
    for line in lines:
        pdf.drawString(x, y, line)
        y -= leading or size * 1.35
    return y


def create():
    fonts()
    data = specs()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(WIDTH), Inches(HEIGHT)
    pdf = canvas.Canvas(str(FINAL / "NUCLE_presentation.pdf"), pagesize=(960, 540))
    pdf.setTitle("NUCLE — kết quả cổ điển và BERT-small")
    notes = ["# Ghi chú thuyết trình\n\n12 slides, khoảng 9–10 phút; chưa có thông tin thời lượng từ người dùng. Tên/MSSV và đóng góp cần nhóm xác nhận.\n"]
    for index, spec in enumerate(data, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        fill = slide.background.fill; fill.solid(); fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(.16))
        bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor.from_string(TEAL); bar.line.fill.background()
        textbox(slide, .55, .42, 12.2, .8, spec["title"], size=30, color=NAVY, bold=True)
        textbox(slide, .55, 7.08, 11.8, .25, "NUCLE • Known-span classification • 28/10/2026", size=10, color=GRAY)
        textbox(slide, 12.2, 7.03, .6, .3, str(index), size=12, color=GRAY)
        pdf.setFillColor(colors.white); pdf.rect(0, 0, 960, 540, fill=1, stroke=0)
        pdf.setFillColor(colors.HexColor("#" + TEAL)); pdf.rect(0, 528, 960, 12, fill=1, stroke=0)
        pdf_text(pdf, 40, 485, spec["title"], size=25, color=NAVY, bold=True, width=880)
        pdf_text(pdf, 40, 18, "NUCLE • Known-span classification • 28/10/2026", size=8, color=GRAY)
        pdf_text(pdf, 905, 18, str(index), size=10, color=GRAY)
        y, py = 1.55, 420
        if "subtitle" in spec:
            textbox(slide, .65, y, 12, .9, spec["subtitle"], size=26, color=TEAL)
            py = pdf_text(pdf, 47, py, spec["subtitle"], size=23, color=TEAL, width=865) - 24
            y += 1.2
        if "image" in spec:
            from PIL import Image
            with Image.open(spec["image"]) as image:
                iw, ih = image.size
            scale = min(11.9 / iw, 3.95 / ih)
            w, h = iw * scale, ih * scale
            slide.shapes.add_picture(str(spec["image"]), Inches((WIDTH-w)/2), Inches(1.25), width=Inches(w), height=Inches(h))
            pdf.drawImage(ImageReader(str(spec["image"])), (960-w*72)/2, 540-(1.25+h)*72, width=w*72, height=h*72, mask='auto')
            y, py = 5.35, 152
        if "table" in spec:
            rows = spec["table"]
            table_shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(.7), Inches(1.45), Inches(11.8), Inches(2.8))
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    cell = table_shape.table.cell(r, c); cell.text = value
                    cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string("DFEDF0" if r == 0 else ("F3F7F9" if r % 2 else "FFFFFF"))
                    for p in cell.text_frame.paragraphs:
                        for run in p.runs: run.font.name = "Arial"; run.font.size = Pt(20); run.font.color.rgb = RGBColor.from_string(INK)
            cell_w, cell_h = 849.6 / len(rows[0]), 40
            for r, row in enumerate(rows):
                yy = 435 - (r + 1) * cell_h
                pdf.setFillColor(colors.HexColor("#DFEDF0" if r == 0 else "#F3F7F9" if r % 2 else "#FFFFFF"))
                pdf.rect(50, yy, 849.6, cell_h, fill=1, stroke=0)
                for c, value in enumerate(row):
                    pdf_text(pdf, 55 + c*cell_w, yy + 15, value, size=15, width=cell_w-10)
            y, py = 4.65, 190
        font_size = 21 if "image" not in spec else 18
        for bullet in spec["bullets"]:
            line_capacity = 84 if font_size == 21 else 98
            lines_estimate = max(1, math.ceil(len(bullet) / line_capacity))
            height = max(.63, lines_estimate * .35 + .12)
            if y + height > 6.92:
                raise ValueError(f"Slide {index} text exceeds allocated height")
            textbox(slide, .72, y, .25, .35, "•", size=font_size, color=ORANGE)
            textbox(slide, 1.03, y, 11.55, height, bullet, size=font_size)
            py = pdf_text(pdf, 72, py, "• " + bullet, size=font_size*.79, width=820) - 16
            y += height + .12
        slide.notes_slide.notes_text_frame.text = spec["notes"]
        notes.append(f"## Slide {index}. {spec['title']} — khoảng {spec['seconds']} giây\n\n{spec['notes']}\n")
        pdf.showPage()
    prs.core_properties.title = "Phân loại lỗi ngữ pháp NUCLE"
    prs.core_properties.author = "Nhóm nghiên cứu — chờ xác nhận tên/MSSV"
    prs.save(FINAL / "NUCLE_presentation.pptx")
    pdf.save()
    (FINAL / "speaker_notes.md").write_text("\n".join(notes), encoding="utf-8")
    dump(FINAL / "presentation_content.json", [{k: str(v.relative_to(ROOT)) if isinstance(v, Path) else v for k, v in row.items()} for row in data])
    print(f"Created {len(data)} slides, PPTX, matching-layout PDF and speaker notes.")


if __name__ == "__main__":
    create()
