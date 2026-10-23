"""Report observed train-only diagnostics without changing prior benchmark scores."""
from pathlib import Path
import json,hashlib,re
from datetime import datetime,timezone
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches,Pt
ROOT=Path(__file__).resolve().parents[2]; OUT=ROOT/'reports/bert_learning_diagnostics'
def read(name): return json.loads((OUT/name).read_text(encoding='utf-8'))
def table(headers,rows):
    return '\n'.join(['| '+' | '.join(headers)+' |','|'+'|'.join(['---']*len(headers))+'|']+['| '+' | '.join(str(c) for c in r)+' |' for r in rows])

def run():
    summary=read('summary.json'); inputs=read('input_audit.json'); gradients=read('gradient_checks.json'); marker=read('marker_sensitivity.json'); pairs=read('same_sentence_pairs.json'); manifest=read('manifest.json')
    assert not summary['validation_scored'] and not summary['test_scored']
    for key,filename in [('protocol_sha256','docs/bert_learning_diagnostics_protocol.md'),('code_sha256','src/evaluation/bert_learning_diagnostics.py'),('readout_code_sha256','src/models/bert_span_readout.py')]:
        assert manifest[key]==hashlib.sha256((ROOT/filename).read_bytes()).hexdigest()
    table_rows=[]; fig,axes=plt.subplots(1,2,figsize=(10,4))
    for run in summary['runs']:
        h=read(run['name']+'/history.json'); first=next((x['updates'] for x in h if x['accuracy']==1.),None)
        table_rows.append([run['config']['mode'],run['config']['learning_rate'],first,run['updates'],f"{run['metrics']['accuracy']:.1%}",f"{run['metrics']['cross_entropy']:.4f}",'Đạt' if run['memorization_passed'] else 'Chưa đạt CE≤0,05'])
        label=run['name']; axes[0].plot([x['updates'] for x in h],[x['cross_entropy'] for x in h],label=label); axes[1].plot([x['updates'] for x in h],[x['accuracy'] for x in h],label=label)
    axes[0].set_yscale('log'); axes[0].set_ylabel('Training cross-entropy'); axes[0].axhline(.05,color='gray',linestyle='--'); axes[1].set_ylabel('Training accuracy'); axes[1].set_ylim(0,1.03)
    for ax in axes: ax.set_xlabel('Optimizer updates'); ax.legend(fontsize=7); ax.grid(alpha=.2)
    fig.suptitle('Memorization diagnostic: same 100 training examples'); fig.tight_layout(); fig.savefig(OUT/'memorization_curves.png',dpi=180); plt.close(fig)
    text='\n\n'.join([
    '# Báo cáo chẩn đoán BERT: khả năng học, gradient và span',
    'Ngày 23/10/2026. Chỉ dùng train. Không tính thêm metric validation/test, không thay kết quả LR hoặc BERT đã công bố. Protocol được ghi trước ba lượt học thuộc; kiểm tra cặp span cùng câu có protocol bổ sung riêng trước khi chạy.',
    '## 1. Kết luận và quyết định',
    '**BERT nhỏ có khả năng học thuộc nhãn của 100 mẫu train; không phát hiện mất gradient trong các kiểm tra đã chạy.** Cả ba cấu hình đạt 100% accuracy trên chính tập đã học. Learning rate cũ 5e-5 đạt 100% accuracy nhưng chưa đạt ngưỡng loss chặt; learning rate 1e-3 đạt cả hai điều kiện nhanh hơn. Vì vậy không có căn cứ nói mô hình hoàn toàn không thể học bài toán.',
    '**Nên thử cấu hình tối ưu hóa khác trên cùng BERT nhỏ, đồng thời so CLS với span mean có kiểm soát.** Chưa cần kết luận phải đổi sang BERT-base hoặc mô hình lớn hơn. Đây là quyết định cho thí nghiệm train/validation tiếp theo, không phải một kết quả validation mới. LR 0,8067 vẫn là ứng viên mạnh nhất theo bằng chứng hiện có; BERT-small cũ 0,1984±0,0030 vẫn là lượt thử chưa thành công.',
    '## 2. Dữ liệu và kiểm tra input',
    f"Kiểm {inputs['counts']['checked']:,} mẫu train: hai marker có ID {inputs['marker_ids']}, khác nhau, là special token và không phải UNK. Không có target bị mất hoặc bị thay bằng UNK; không có input tokenized trùng hoặc trùng nhưng khác nhãn. Có {inputs['counts']['truncated']} câu bị cắt context, span vẫn giữ nguyên. Chỉ lấy sample sau khi kiểm partition=train; không tạo model input cho validation/test.",
    'Chọn 100 mẫu bằng seed20261023, 20 mẫu mỗi nhãn, không dựa trên prediction. Các lượt dùng cùng input/labels, cùng thứ tự minibatch và cùng hash trọng số khởi tạo. Đây là tập cân bằng có chủ đích, dễ hơn bài toán đánh giá tổng quát hóa; có thể học cả danh tính câu, nên học thuộc không chứng minh model đã hiểu ngữ pháp hoặc dùng span đúng cách.',
    '## 3. Gradient và cập nhật tham số',
    'Đo trên một minibatch train cân bằng 20 mẫu, cả pretrained model và checkpoint CLS thất bại trước đó. Dropout tắt trong phép probe để kiểm tra xác định. Loss khớp cross-entropy tính độc lập; forward của wrapper CLS khớp forward Hugging Face gốc. Optimizer chứa đủ tham số trainable, không trùng lặp. Không có gradient thiếu hoặc NaN/Inf.',
    table(['Nguồn / readout','Embedding norm','Encoder layer0 norm','Layer1 norm','Classifier norm','Δ marker đầu sau1step'],[[k,f"{v['gradient_norms']['embeddings']:.6f}",f"{v['gradient_norms']['encoder.layer.0']:.6f}",f"{v['gradient_norms']['encoder.layer.1']:.6f}",f"{v['gradient_norms']['classifier']:.6f}",f"{v['parameter_update_norms']['start_marker']:.6f}"] for k,v in gradients.items()]),
    'Embedding của cả hai marker, trọng số query ở encoder và classifier đều thay đổi sau một bước AdamW. Điều này loại trừ dấu hiệu đóng băng/đứt gradient tại các thành phần đã kiểm, nhưng không chứng minh mọi chi tiết huấn luyện đều tối ưu hoặc không thể có lỗi khác.',
    '### Vì sao gradient tại hidden target cuối bằng 0 với CLS?',
    'Readout cũ lấy hCLS từ layer cuối rồi qua dense+tanh và classifier. Vì head không đọc trực tiếp hidden target cuối, gradient trực tiếp tại vị trí target của layer cuối bằng0; gradient ở target của các layer trước và embedding vẫn khác0 qua self-attention. Đây là cấu trúc tính toán dự kiến, không phải bằng chứng encoder không học từ span.',
    'Span mean lấy hspan=(Σt∈S ht)/|S|, S là các subword giữa marker, không gồm marker hoặc padding. Giữ cùng dense+tanh pooler và classifier; không thêm tham số. Khi đó gradient trực tiếp đi vào hidden target cuối. Span rỗng có fallback marker đầu trong code và test; dữ liệu train hiện không có trường hợp này.',
    '## 4. Mô hình có phản ứng với vị trí span không?',
    'Trên 100 mẫu train, giữ thứ tự các token nội dung và dời cặp marker sang một vùng khác không giao span cũ, có cùng số token. Dùng checkpoint CLS đã huấn luyện trước đây. Phép perturbation này đổi cả vị trí tương đối của marker/token; không phải câu có ground truth mới để chấm accuracy.',
    table(['Readout dùng với encoder cũ','Số prediction thay đổi','Mean L1 xác suất','Cosine representation trung bình'],[[m,f"{marker[m]['prediction_changes']}/100",f"{marker[m]['mean_probability_L1']:.4f}",f"{marker[m]['mean_representation_cosine']:.4f}"] for m in ['cls','span_mean']]),
    f"Pooled CLS đổi prediction ở {marker['cls']['prediction_changes']}/100 mẫu: mô hình không hoàn toàn bỏ qua việc thay marker. Tuy nhiên phản ứng với perturbation không chứng minh nó định vị loại lỗi chính xác. Span mean trong bảng này chỉ thay cách đọc encoder cũ qua cùng classifier chưa được fit cho mean; không dùng nó để so chất lượng phân loại. Thêm padding làm logits thay đổi tối đa {marker['padding_max_logit_difference']:.2e}, phù hợp sai số số học nhỏ.",
    '### Hai gold span khác nhãn trong cùng câu',
    f"Từ {pairs['candidate_sentence_count']} câu train đủ điều kiện, chọn50 câu có hai gold span khác vị trí và khác nhãn. Checkpoint CLS cũ cho prediction khác nhau ở {pairs['predictions_differ']}/50 cặp; chỉ {pairs['both_correct']}/50 cặp đúng cả hai. Đây là stress test trên train, chọn có chủ đích, không phải validation accuracy. Kết quả cho thấy phản ứng với marker chưa đồng nghĩa phân biệt tốt hai mục tiêu thật trong cùng câu.",
    f"Cosine trung bình giữa hai biểu diễn CLS là {pairs['mean_cls_cosine']:.4f}; giữa hai mean-span là {pairs['mean_span_mean_cosine']:.4f}. Mean-span thay đổi nhiều hơn theo mục tiêu là điều có thể mong đợi từ cách đọc; số này không chứng minh mean-span sẽ có F1 cao hơn sau huấn luyện.",
    '## 5. Học thuộc 100 mẫu train',
    'Mỗi lượt khởi tạo lại cùng pretrained BERT/seed, không tiếp tục từ checkpoint thất bại. Fine-tune toàn bộ encoder và head; dropout0, weight decay0, learning rate cố định, batch20, clip norm1, tối đa400 update. Mỗi25 update đánh giá lại chính100 mẫu. Ngưỡng chặt: accuracy=1 và CE≤0,05 ở hai lần kiểm liên tiếp. Đây là phép thử khả năng tối ưu hóa trong điều kiện thuận lợi, không tái lập đầy đủ recipe huấn luyện cũ.',
    table(['Readout','Learning rate','Lần đầu quan sát acc100%','Update cuối','Accuracy cuối','CE cuối','Tiêu chí chặt'],table_rows),
    '![Đường học thuộc](../reports/bert_learning_diagnostics/memorization_curves.png)',
    'CLS5e-5 đã dự đoán đúng cả100 mẫu từ mốc175 update; loss vẫn0,1515 ở update400 nên chưa qua tiêu chí confidence chặt đã đặt trước. Không được gọi kết quả này là “không học thuộc được nhãn”. CLS1e-3 qua tiêu chí ở update100; span mean1e-3 ở update75. Với batch20, mỗi5 update tương ứng một lượt qua100 mẫu; số lần thấy lại mỗi ví dụ cao hơn nhiều so với thí nghiệm toàn train5epoch trước đó.',
    'Trong cùng phép thử100 mẫu, khác biệt giữa hai lượt CLS là learning rate; span mean so với CLS ở1e-3 chỉ đổi readout qua cùng pooler/head. Tuy nhiên khác biệt giữa phép thử nhỏ và lượt thất bại trên toàn train còn gồm số lần lặp mỗi ví dụ, dropout, weight decay, schedule và phân bố lớp. Không quy toàn bộ thất bại cũ cho learning rate chỉ từ kết quả này.',
    '## 6. Quyết định sau chẩn đoán',
    '1. Giữ kết luận BERT cũ là kết quả chưa thành công của cấu hình đã thử; không suy ra ngữ cảnh vô ích. LR vẫn là ứng viên mạnh nhất trên validation.',
    '2. Có căn cứ thử lại tối ưu hóa trên cùng checkpoint nhỏ. Learning rate1e-3 ở tập học thuộc là giá trị chẩn đoán, chưa phải hyperparameter đã chọn cho toàn bộ dữ liệu.',
    '3. Nên so CLS và span mean với cùng split, ngân sách, optimizer và cách chọn checkpoint trên validation trong protocol mới. Giữ dropout/regularization cần thiết cho generalization; không bê nguyên recipe dropout0 của phép học thuộc sang benchmark.',
    '4. Trước khi tăng kích thước mô hình, ưu tiên xem learning curve và khả năng phân biệt hai target trong cùng câu. Span mean là ứng viên hợp lý để thử, chưa có bằng chứng nó vượt LR.',
    '5. Lần này dừng ở chẩn đoán và quyết định; chưa chạy grid mới trên toàn train/validation, chưa chấm test hoặc khóa mô hình cuối.',
    '## 7. Code, artifact và kiểm chứng',
    table(['File','Vai trò'],[
    ['src/models/bert_span_readout.py','target_mask và wrapper CLS/span mean, chia sẻ encoder/pooler/head'],
    ['src/evaluation/bert_learning_diagnostics.py','Train-only loader, chọn100 mẫu, gradient probe, marker perturbation và ba lượt học thuộc'],
    ['src/evaluation/bert_target_pair_probe.py','Stress test50 cặp gold span khác nhãn cùng câu train'],
    ['tests/test_bert_readout.py','Mask, insertion fallback, marker lỗi, forward CLS tương đương, gradient target và padding'],
    ['src/reporting/build_bert_diagnostics_report.py','Đọc artifact và tạo báo cáo/đồ thị']]),
    '13 kiểm tra đã qua:7 kiểm tra readout mới và6 kiểm tra tokenizer/span cũ. Cả ba checkpoint chẩn đoán tái tạo đúng prediction train đã lưu. Với span mean phải tải BertReadout(mode=span_mean), không tải head CLS mặc định rồi coi là cùng mô hình; mode nằm trong config.json mỗi lượt.',
    'Artifact local: reports/bert_learning_diagnostics/ gồm input_audit.json, selected_train_ids.csv, gradient_checks.json, marker_sensitivity.json, same_sentence_pairs.json, history/checkpoint/predictions từng lượt và summary.json. Không ghi câu corpus hoặc correction vào report. Hash code/protocol/split/samples và environment được ghi ở manifest.json. Source pretrained và checkpoint cũ chỉ đọc, không ghi đè.',
    'Công thức loss: CE=−(1/N)Σi log softmax(zi)[yi]. Gradient norm nhóm g là sqrt(Σp∈g Σj (∂CE/∂pj)²). Cosine(a,b)=a·b/(||a||||b||); L1 xác suất=Σk|pk−p′k|. Các metric học thuộc chỉ tính trên100 mẫu train đã fit, không so trực tiếp với macro-F1 validation của LR.'
    ])+'\n'
    path=ROOT/'docs/bert_learning_diagnostics_report.md'; path.write_text(text,encoding='utf-8')
    doc=Document(); doc.styles['Normal'].font.name='Calibri'; doc.styles['Normal'].font.size=Pt(10)
    current=None
    for line in text.splitlines():
        if line.startswith('|'):
            cells=[x.strip() for x in line.strip('|').split('|')]
            if all(re.fullmatch('[-:]+',x) for x in cells): continue
            if current is None: current=doc.add_table(rows=0,cols=len(cells)); current.style='Table Grid'
            row=current.add_row()
            for cell,value in zip(row.cells,cells): cell.text=value
            continue
        current=None
        if not line: continue
        if line.startswith('!['): doc.add_picture(str(OUT/'memorization_curves.png'),width=Inches(6)); continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#')); doc.add_heading(line[level:].strip(),level=min(level,3))
        else: doc.add_paragraph(line.replace('**',''))
    doc.save(ROOT/'docs/bert_learning_diagnostics_report.docx')
    inventory={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in OUT.rglob('*') if p.is_file() and p.name!='artifact_hashes.json'}
    (OUT/'artifact_hashes.json').write_text(json.dumps(inventory,indent=2),encoding='utf-8')
    print('Diagnostics report generated.')
if __name__=='__main__': run()
