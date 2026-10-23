"""Build week-1 validation metrics and Markdown/DOCX report from saved artifacts."""
from pathlib import Path
import csv, json, sys, statistics, re, hashlib
from collections import Counter
ROOT=Path(__file__).resolve().parents[2]; sys.path.insert(0,str(ROOT))
from src.evaluation.nucle_span_analysis import metric, LABELS, PAIRS
from src.evaluation.metrics import scores
OUT=ROOT/'reports/contextual_model_validation'
def readj(p): return json.loads(p.read_text(encoding='utf-8'))
def readcsv(p):
    with p.open(encoding='utf-8-sig',newline='') as f: return list(csv.DictReader(f))
def dump(p,o): p.write_text(json.dumps(o,ensure_ascii=False,indent=2),encoding='utf-8')
def fmt(v): return 'N/A' if v is None else f'{v:.4f}'
def table(headers,rows):
    return '\n'.join(['| '+' | '.join(headers)+' |','|'+'|'.join(['---']*len(headers))+'|']+['| '+' | '.join(str(v) for v in row)+' |' for row in rows])

def run():
    selection=readj(OUT/'bert/selection.json'); audit=readj(OUT/'audit/summary.json'); reviews=readj(OUT/'audit/manual_review.json')
    base=readcsv(ROOT/'reports/nucle_feature_benchmark_solver_v2/validation_predictions.csv')
    val=[r['sample_id'] for r in base]; samples={}; meta={}
    with (ROOT/'data_set/processed/nucle_preparation_v2/samples.jsonl').open(encoding='utf-8') as f:
        for line in f:
            r=json.loads(line)
            if r['sample_id'] in set(val):
                samples[r['sample_id']]={'error_type':r['error_type']}
                meta[r['sample_id']]={'tokens':len(r['span_token_indices'])}
    posrows=readcsv(ROOT/'reports/nucle_pos_baseline_v1/validation_predictions.csv'); pos={r['sample_id']:r['pos_signature'] for r in posrows}
    for row in base+posrows:
        assert row['error_type']==samples[row['sample_id']]['error_type']
    predictions={name:{r['sample_id']:r[key] for r in base} for name,key in [('Majority','majority'),('LR Full (saved)','best_lr'),('Tree Full','best_tree')]}
    predictions['POS rule']={r['sample_id']:r['prediction'] for r in posrows}
    ab=readcsv(ROOT/'reports/nucle_feature_ablation_v1/validation_predictions.csv')
    predictions['Tree S+L+P+M']={r['sample_id']:r['prediction'] for r in ab if r['model']=='DecisionTree' and r['variant']=='S+L+P+M'}
    bert_names=[]
    for run in selection['runs']:
        if run['lr']!=selection['selected_lr']: continue
        name=f"BERT-small seed {run['seed']}"; bert_names.append(name)
        bert_rows=readcsv(ROOT/run['directory']/'validation_predictions.csv')
        for row in bert_rows: assert row['error_type']==samples[row['sample_id']]['error_type']
        predictions[name]={r['sample_id']:r['prediction'] for r in bert_rows}
    primary=f"BERT-small seed {selection['primary_seed']}"
    slices={'all':(val,LABELS),'empty':([s for s in val if meta[s]['tokens']==0],LABELS),'single_token':([s for s in val if meta[s]['tokens']==1],LABELS),'multi_token':([s for s in val if meta[s]['tokens']>1],LABELS)}
    for pair,labels in PAIRS.items():
        ids=[s for s in val if samples[s]['error_type'] in labels]; slices[pair]=(ids,labels)
        for signature in ('VERB','AUX','NOUN','DET'): slices[pair+' POS='+signature]=([s for s in ids if pos[s]==signature],labels)
    results={}; flat=[]; classrows=[]
    for name,pred in predictions.items():
        assert set(pred)==set(val) and set(pred.values())<=set(LABELS)
        results[name]={}
        for slice_name,(ids,labels) in slices.items():
            m=metric(ids,samples,pred,labels)
            if ids:
                m['macro_precision']=sum(v['precision'] for v in m['per_label'].values())/len(labels)
                m['macro_recall']=sum(v['recall'] for v in m['per_label'].values())/len(labels)
            results[name][slice_name]=m
            flat.append({'model':name,'slice':slice_name,**{k:m[k] for k in ('support','macro_f1','accuracy','outside_label_predictions')}})
            for label,c in m['per_label'].items(): classrows.append({'model':name,'slice':slice_name,'label':label,**c})
        results[name]['confusion_matrix']=scores([samples[s]['error_type'] for s in val],[pred[s] for s in val])['confusion_matrix_true_rows_predicted_columns']
        results[name]['prediction_counts']=dict(Counter(pred.values()))
    for filename,rows in [('comparison.csv',flat),('per_class_metrics.csv',classrows)]:
        with (OUT/filename).open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.DictWriter(f,fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
    seed_aggregate={}
    for slice_name in slices:
        vals=[results[name][slice_name]['macro_f1'] for name in bert_names]
        seed_aggregate[slice_name]={'mean':statistics.mean(vals) if vals[0] is not None else None,'sample_std':statistics.stdev(vals) if vals[0] is not None else None}
    result={'test_scored':False,'primary_bert':primary,'models':results,'bert_seed_aggregate':seed_aggregate,'input_sha256':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [OUT/'bert/selection.json',OUT/'audit/summary.json',ROOT/'reports/nucle_feature_benchmark_solver_v2/validation_predictions.csv']}}
    dump(OUT/'comparison.json',result)
    # Save standalone figures; no corpus text.
    import matplotlib; matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    names=['Majority','POS rule','LR Full (saved)','Tree S+L+P+M',primary]
    fig,ax=plt.subplots(figsize=(8,4)); bars=ax.bar(['Majority','POS','LR full','Tree best','BERT small'],[results[n]['all']['macro_f1'] for n in names])
    ax.set_ylim(0,1); ax.set_ylabel('Validation macro-F1'); ax.bar_label(bars,fmt='%.3f'); fig.tight_layout(); fig.savefig(OUT/'overall_macro_f1.png',dpi=180); plt.close(fig)
    fig,axes=plt.subplots(1,2,figsize=(10,4))
    for ax,name in zip(axes,['LR Full (saved)',primary]):
        matrix=results[name]['confusion_matrix']; ax.imshow(matrix,cmap='Blues'); ax.set_xticks(range(5),LABELS,rotation=45); ax.set_yticks(range(5),LABELS); ax.set_title(name); ax.set_xlabel('Predicted'); ax.set_ylabel('Gold')
        for i in range(5):
            for j in range(5): ax.text(j,i,str(matrix[i][j]),ha='center',va='center',color='red',fontsize=8)
    fig.tight_layout(); fig.savefig(OUT/'confusion_matrices.png',dpi=180); plt.close(fig)
    selected=results[primary]; lr=results['LR Full (saved)']; delta=selected['all']['macro_f1']-lr['all']['macro_f1']
    lines=['# Báo cáo kết quả tuần 1 — NUCLE','', 'Ngày thực hiện: 23/10/2026. Các task được thực hiện trong phiên làm việc này theo thứ tự của kế hoạch tuần 1; không giả định đã trôi qua bảy ngày. Phạm vi chỉ train/validation; không sửa Excel, không khóa mô hình cuối, không chấm test, không bootstrap hoặc làm các task tuần 2.','',
    '## 1. Kết quả chính','',
    f"LR Full đã lưu đạt macro-F1 {fmt(lr['all']['macro_f1'])}; BERT-small seed chính đạt {fmt(selected['all']['macro_f1'])}, chênh lệch BERT−LR {delta:+.4f}. BERT dùng kiến trúc nhỏ 2 layer/hidden 128 trên CPU, không phải BERT-base. Kết quả ba seed tại learning rate được chọn: {fmt(selection['selected_lr_macro_f1_mean'])} ± {fmt(selection['selected_lr_macro_f1_sample_std'])} (sample standard deviation, không phải confidence interval).",'',
    '## 2. Hoàn thành từng task','',table(['Task tuần 1','Việc đã làm','Bằng chứng'],[
    ['Related work','Bảng input/output, dữ liệu, metric và giới hạn áp dụng','docs/contextual_models_related_work.md; reference.md'],
    ['Audit validation','Kiểm offset/alignment train+validation; đọc 18 lỗi có chủ đích và đối chiếu correction sau dự đoán','audit/summary.json; manual_review.json'],
    ['Sensitivity v1/v2','Cùng split, LR C=0.1, no class weight; encoder fit riêng trên train','audit/v1_validation_predictions.csv; v2_validation_predictions.csv'],
    ['Protocol và pilot BERT','Checkpoint pin revision, marker, crop giữ span, ngân sách sau pilot','docs/bert_experiment_protocol.md; bert/pilot.json'],
    ['Fine-tuning','Hai learning rate ở seed đầu; learning rate được chọn chạy thêm hai seed; 5 epoch mỗi lượt','bert/*/history.json, checkpoint, summary.json'],
    ['So sánh','Toàn tập, từng lớp, span, cặp nhãn, POS slices; giữ Majority/Tree','comparison.csv; per_class_metrics.csv; comparison.json']]),'',
    '## 3. Related work và quyết định thiết kế','',
    'Bảng literature đầy đủ nằm ở docs/contextual_models_related_work.md; danh mục nguồn ở reference.md. REF11 hỗ trợ thử encoder ngữ cảnh; REF07 giải thích vai trò loại lỗi trong GEC; REF12 hỗ trợ đánh giá theo hiện tượng; REF13/REF15 hỗ trợ thảo luận đặc trưng/parser. Không dùng điểm các bài ngoài làm đối chứng trực tiếp. REF14 và REF20 mô tả BERT và mô hình pretrained nhỏ; mô hình span-marked năm lớp là điều chỉnh của nhóm. Không triển khai LET/XGEC.', '',
    '## 4. Audit dữ liệu và sensitivity','',
    'Giữ 11.052 train và 2.368 validation. Không giữ/chấm test trong code thực nghiệm mới. V1/v2 có cùng ID và nhãn; 381 mẫu train và 96 validation thay câu hoặc offset nội câu. Kiểm tra offset khớp span và số token giao span cho mọi mẫu train/validation: không phát hiện bất đồng. Không có exact input signature xuất hiện đồng thời ở train và validation; chưa kiểm chứng mọi near-duplicate ngữ nghĩa.','',
    table(['Phiên bản','Chiều feature','POS macro-F1','LR macro-F1','LR: 96 mẫu đổi input','LR: mẫu không đổi'],[[v,audit[v]['dimension'],fmt(audit[v]['pos_macro']['macro_f1']),fmt(audit[v]['macro']['macro_f1']),fmt(audit[v]['changed']['macro_f1']),fmt(audit[v]['unchanged']['macro_f1'])] for v in ['v1','v2']]),'',
    'Sensitivity giữ cấu hình, không retune. Cả hai LR dừng theo change tolerance trước giới hạn, không phải gradient tolerance. Điểm trên tập con tính lại theo năm lớp, không phải phần đóng góp cộng được vào điểm chung. Vì train cũng đổi, thay đổi prediction trên mẫu validation không đổi vẫn có thể xảy ra. Chênh lệch nhỏ chưa có CI và không chứng minh v2 cải thiện mô hình.', '',
    f"Refit v2 khác baseline đã lưu ở {audit['v2_reproduction_prediction_disagreements']}/2.368 dự đoán. Chẩn đoán giữ nguyên feature matrix: 4 threads tái tạo prediction sensitivity, 10 threads tái tạo prediction baseline cũ (0 bất đồng ở mỗi đối chiếu tương ứng). LBFGS dừng theo change tolerance ở iteration khác nhau. Đây là bằng chứng độ nhạy số học theo threads trong lần kiểm tra; bảng BERT giữ baseline gốc, không thay bằng điểm refit cao hơn. Chi tiết: audit/thread_diagnostic.json.", '',
    'Alignment validation v2: '+str(audit['v2']['alignment_counts_validation'])+'. Partial-token không tự là lỗi kỹ thuật: contract lấy token giao span. Kiểm cơ học không chứng minh parser đúng ngữ nghĩa.','',
    '### Audit 18 lỗi tiêu biểu','',
    'Mẫu được chọn có chủ đích: ba ID đầu theo thứ tự cho mỗi chiều Vform/SVA, Nn/ArtOrDet; thêm sáu lỗi nhiều token. Đây là review hỗ trợ bằng AI, không phải adjudication độc lập của chuyên gia và không ước lượng tỷ lệ nguyên nhân trong corpus. Không đổi gold label. Correction chỉ đọc sau prediction để audit; không đưa vào predictor.','',
    table(['Annotation ID','Nhóm nhận xét','Mức chắc chắn','Giải thích'],[[r['annotation_id'],r['category'],r['confidence'],r['note']] for r in reviews['rows']]),'',
    'Điểm mới quan trọng: lỗi thiếu determiner có thể được annotation bằng span danh từ không rỗng. Vì vậy support span rỗng bằng 0 không đồng nghĩa corpus không chứa lỗi thiếu từ. Một số span có nhiều thay đổi trong correction hoặc gold label chồng lấp với loại ngữ pháp khác, nên không quy mọi lỗi phân loại cho mô hình.', '',
    '### Khác biệt code–protocol đã làm rõ','',
    'Baseline thực tế có vocabulary tối đa 500 ở kênh pooled n-gram, cộng categorical vocabulary riêng cho slot lexical trái/phải, không phải tổng vocabulary chung 500. Numeric chưa z-score. Feature parse lại câu gốc; masking xóa lexical span và fit encoder riêng nhưng giữ parse. Những khác biệt này được ghi ở docs/implementation_data_audit.md và mục 9 protocol; không thay baseline trong lần này.', '',
    '## 5. BERT: input, công thức và huấn luyện','',
    'Input là câu gốc có [SPAN_START]/[SPAN_END] quanh gold span, thêm [CLS]/[SEP]. Tokenize trái/span/phải riêng; max length 128, ưu tiên giữ toàn bộ span rồi lấy ngữ cảnh gần hai phía. Marker mới được thêm embedding; classifier và toàn encoder cùng được fine-tune. Không có nhãn, correction hoặc ID trong input model.','',
    'Công thức: h = BERTθ(input)[pooled CLS]; z = W·Dropout(h)+b; p(y=k|x)=exp(zₖ)/Σⱼexp(zⱼ); LCE=−(1/N)Σᵢlog p(yᵢ|xᵢ). AdamW dùng weight decay tách khỏi gradient cross-entropy, không mô tả nó như LR L2 objective.','',
    'F1ₖ=2TPₖ/(2TPₖ+FPₖ+FNₖ), zero division=0; macro-F1=(1/5)ΣₖF1ₖ. Trên gold-label pair H, tính lại F1 từng nhãn trong H rồi lấy trung bình hai nhãn; prediction ngoài cặp vẫn gây FN. Không ép argmax về hai lớp. Độ lệch chuẩn giữa seed: s=√[Σᵣ(F1ᵣ−mean)²/(R−1)].','',
    'Checkpoint google/bert_uncased_L-2_H-128_A-2, revision 30b0a37ccaaa32f332884b96992754e246e48c5f. Transformers 4.57.6, PyTorch 2.14.0+cpu; 4 threads; khoảng 4,39 triệu tham số. Pilot 20 bước ~2,314 giây. Có 5 train/6 validation bị cắt ngữ cảnh, không span bị cắt. Tokenize quanh ranh giới span có thể chia lại một từ thành subword khác; đây là lựa chọn cần ghi nhận.','',
    'Batch 32, 5 epoch, AdamW, warmup 10% rồi linear decay, weight decay 0,01 trừ bias/LayerNorm, clip norm 1. Grid learning rate 3e-5 và 5e-5 ở seed 20261013; chọn checkpoint bằng macro-F1 validation, hòa chọn epoch sớm nhất; LR hòa chọn đầu grid. LR chọn chạy thêm seed 20261014/20261015; không chọn seed tốt nhất đại diện.','',
    table(['Learning rate','Seed','Epoch được chọn','Macro-F1','Runtime giây'],[[r['lr'],r['seed'],r['best_epoch'],fmt(r['metrics']['macro_f1']),round(r['seconds'],1)] for r in selection['runs']]),'',
    '### Code và artifact','',
    table(['File/hàm','Vai trò'],[
    ['src/evaluation/preprocessing_sensitivity_audit.py::run','Chỉ lấy train/validation; kiểm input; parse và fit encoder riêng mỗi preparation; fit LR cố định; xuất packet audit'],
    ['src/training/bert_span_classifier.py::encode_span','Kiểm offset, tokenize ba vùng và crop giữ đủ span'],
    ['load_data / make_loader','Loại test trước khi tạo model input; nhãn riêng; dynamic padding'],
    ['new_model / run / train_one','Tải checkpoint local; init classifier/marker; reset seed và optimizer từng run; chọn checkpoint validation'],
    ['src/reporting/build_contextual_validation_report.py','Đọc prediction lưu, kiểm IDs/labels, tính metrics/slices, biểu đồ và report'],
    ['tests/test_bert_span_classifier.py','Kiểm crop, partial/empty span, bounds sai và span vượt budget']]),'',
    '## 6. Kết quả validation','',table(['Mô hình','Support','Macro-F1','Accuracy','Macro precision','Macro recall'],[[name,results[name]['all']['support'],fmt(results[name]['all']['macro_f1']),fmt(results[name]['all']['accuracy']),fmt(results[name]['all']['macro_precision']),fmt(results[name]['all']['macro_recall'])] for name in names+bert_names[1:]]),'',
    '![Macro-F1](../reports/contextual_model_validation/overall_macro_f1.png)','',
    'Tree best ở bảng là variant S+L+P+M được chọn từ validation trước đây; Tree Full cũng được lưu trong comparison.csv. Các seed bổ sung dùng LR được chọn từ seed đầu; mean/std không phải ba lượt chọn hyperparameter độc lập.','',
    '### Từng lớp: BERT seed chính và LR Full','',table(['Nhãn','Support','LR F1','BERT precision','BERT recall','BERT F1'],[[lab,lr['all']['per_label'][lab]['support'],fmt(lr['all']['per_label'][lab]['f1']),fmt(selected['all']['per_label'][lab]['precision']),fmt(selected['all']['per_label'][lab]['recall']),fmt(selected['all']['per_label'][lab]['f1'])] for lab in LABELS]),'',
    '### Span, cặp nhãn và POS slices','',table(['Nhóm','n','POS F1','LR F1','BERT seed chính F1','BERT mean ± std'],[[s,selected[s]['support'],fmt(results['POS rule'][s]['macro_f1']),fmt(lr[s]['macro_f1']),fmt(selected[s]['macro_f1']),fmt(seed_aggregate[s]['mean'])+' ± '+fmt(seed_aggregate[s]['sample_std'])] for s in slices if s!='all']),'',
    'Span bins dùng macro-F1 năm lớp; pair/POS slices dùng macro-F1 hai nhãn trong cặp. Giữ đủ nhãn ngay cả khi support một lớp bằng 0, vì vậy nhóm rất nhỏ phải đọc cùng support và per_class_metrics.csv. Span rỗng support=0 ghi N/A. Không gọi mọi mẫu cặp nhãn là chắc chắn khó và không coi POS NOUN là noun phrase đã xác minh.','',
    '![Confusion matrices](../reports/contextual_model_validation/confusion_matrices.png)','',
    'Phân bố prediction BERT seed chính: '+str(results[primary]['prediction_counts'])+'. Các ô confusion đầy đủ nằm trong comparison.json; hàng là gold, cột là prediction theo thứ tự năm nhãn.','',
    '## 7. Diễn giải, hạn chế và ứng viên bàn giao','',
    f"Chênh lệch BERT−LR toàn tập là {delta:+.4f}; trên Vform/SVA là {selected['Vform/SVA']['macro_f1']-lr['Vform/SVA']['macro_f1']:+.4f}, trên Nn/ArtOrDet là {selected['Nn/ArtOrDet']['macro_f1']-lr['Nn/ArtOrDet']['macro_f1']:+.4f}. Kết luận phải giới hạn ở checkpoint nhỏ, pooled CLS, marker mới và ngân sách năm epoch này. Không ngoại suy thành BERT-base hoặc mọi contextual model.", '',
    'BERT seed chính tập trung dự đoán ArtOrDet/Nn và không dự đoán Vform/SVA trong lượt này. Accuracy cao hơn Majority một ít không đồng nghĩa đã học đủ năm loại lỗi; macro-F1 và F1 từng lớp cho thấy hạn chế rõ. Các giả thuyết gồm dung lượng nhỏ, pooled CLS chưa tập trung vào span và ngân sách tối ưu hóa hạn chế; chưa có ablation để xác nhận nguyên nhân. Không kéo dài tuning sau khi xem kết quả trong lượt đã chốt. Checkpoint replay và train metrics giúp kiểm tra lỗi lưu/tải và khả năng fit, nhưng không chứng minh nguyên nhân ngôn ngữ. Span pooling hoặc BERT-base là hướng thử sau này, chưa thực hiện trong tuần 1.', '',
    'Ứng viên chuyển sang bước lựa chọn của tuần 2: LR Full đã hội tụ theo contract, Tree S+L+P+M, POS/Majority làm đối chứng; BERT-small được giữ như kết quả thí nghiệm ngữ cảnh và có checkpoint/seed/config truy vết. Đây không phải quyết định khóa mô hình cuối. Sensitivity refit không được dùng để thay baseline hoặc chọn điểm cao nhất.', '',
    'Chưa có test/CI; validation đã dùng cho phát triển nên có selection bias. Audit 18 mẫu có chủ đích và do AI hỗ trợ, không đại diện toàn corpus. Pretraining BERT có dữ liệu ngoài train NUCLE, được phân biệt với quy tắc supervised fit chỉ trên train; chưa kiểm chứng pretraining contamination. Không so trực tiếp với số liệu các bài ngoài.', '',
    '## 8. Kiểm tra và tái tạo trong môi trường hiện tại','',
    'Đã qua 6 kiểm tra mới về span/tokenization và không dùng nhãn/correction làm feature, cùng 12 kiểm tra hiện có về feature, benchmark, alignment và boundary. Đây là kiểm tra chức năng trong môi trường hiện tại, không phải task tái lập môi trường sạch của tuần 2. Environment freeze: reports/contextual_model_validation/environment.txt; thư viện bổ sung: requirements-bert.txt.', '',
    'Lệnh lần chạy: .venv/Scripts/python.exe src/evaluation/preprocessing_sensitivity_audit.py; .venv/Scripts/python.exe src/training/bert_span_classifier.py --pilot; .venv/Scripts/python.exe src/training/bert_span_classifier.py --train; .venv/Scripts/python.exe src/reporting/build_contextual_validation_report.py. Audit/train từ chối ghi đè thư mục lượt chạy; sao lưu hoặc chọn output mới khi tái tạo, không xóa artifact chuẩn.', '',
    'Dữ liệu/packet/checkpoint/predictions lưu local dưới data_set hoặc reports đã bị Git ignore. Report không chứa câu corpus/correction. Hash đầu vào, code, protocol nằm trong bert/manifest.json và audit/summary.json. Không thực hiện task tuần 2.']
    md='\n'.join(lines)+'\n'
    replay=OUT/'bert/checkpoint_replay.json'
    if replay.exists(): md+='\n### Kiểm checkpoint đã lưu\n\n'+table(['Run','Prediction validation khác bản lưu','Train macro-F1','Train accuracy'],[[r['run'],r['prediction_disagreements'],fmt(r['train']['macro_f1']),fmt(r['train']['accuracy'])] for r in readj(replay)])+'\n'
    path=ROOT/'docs/contextual_model_validation_report.md'; path.write_text(md,encoding='utf-8')
    from docx import Document
    from docx.shared import Inches, Pt
    doc=Document(); doc.styles['Normal'].font.name='Calibri'; doc.styles['Normal'].font.size=Pt(10)
    in_table=False; current=None
    for line in md.splitlines():
        if line.startswith('|'):
            cells=[c.strip() for c in line.strip('|').split('|')]
            if all(re.fullmatch(r'[-:]+',c) for c in cells): continue
            if not in_table: current=doc.add_table(rows=0,cols=len(cells)); current.style='Table Grid'; in_table=True
            row=current.add_row()
            for cell,text in zip(row.cells,cells): cell.text=text
            continue
        in_table=False
        if not line.strip(): continue
        if line.startswith('!['):
            image=re.search(r'\]\((.*?)\)',line).group(1); doc.add_picture(str((path.parent/image).resolve()),width=Inches(6)); continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#')); doc.add_heading(line[level:].strip(),level=min(level,3))
        else: doc.add_paragraph(line)
    doc.save(ROOT/'docs/contextual_model_validation_report.docx')
    print('Report generated; BERT primary',selected['all']['macro_f1'],'LR',lr['all']['macro_f1'])
if __name__=='__main__': run()
