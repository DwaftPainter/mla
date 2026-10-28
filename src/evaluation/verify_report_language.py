"""Check that a report language edit preserved tables, examples and evidence."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import re

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
REPORTS = ROOT / "reports"
BACKUP = REPORTS / "documentation_archive/before_language_edit_20261028"


def sha(path):
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def table_hash(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False).encode()).hexdigest()


def snippets(text):
    result, current, marker = [], [], None
    for line in text.splitlines():
        if line.strip().startswith(("```", "~~~")):
            if marker is None:
                marker = line.strip()[:3]; current = [line]
            elif line.strip().startswith(marker):
                result.append("\n".join(current + [line])); current = []; marker = None
            continue
        if marker is not None:
            current.append(line)
        elif line.startswith(("    ", "\t")):
            result.append(line)
    return result


def run():
    local = ROOT / ".cache/report_editorial"
    inventory = json.loads((local / "inventory.json").read_text(encoding="utf-8"))
    protected = json.loads((local / "protected.json").read_text(encoding="utf-8"))
    changed_evidence = [name for name, expected in protected.items() if not (ROOT / name).exists() or sha(ROOT / name) != expected]
    checks = []
    edited_paragraphs = {}
    markdown_changed, word_changed = 0, 0
    for name, item in inventory.items():
        path = REPORTS / name
        original = BACKUP / name
        if path.suffix == ".md":
            before, after = original.read_text(encoding="utf-8-sig"), path.read_text(encoding="utf-8-sig")
            tables_before = re.findall(r"(?:^\|[^\n]*\n)+", before, re.MULTILINE)
            tables_after = re.findall(r"(?:^\|[^\n]*\n)+", after, re.MULTILINE)
            checks.append({"file": name, "tables_unchanged": tables_before == tables_after,
                           "table_sha256_before": table_hash(tables_before), "table_sha256_after": table_hash(tables_after),
                           "executable_examples_unchanged": snippets(before) == snippets(after)})
            markdown_changed += before != after
        elif path.suffix == ".docx":
            doc = Document(path)
            values = [[[c.text for c in row.cells] for row in table.rows] for table in doc.tables]
            checks.append({"file": name, "tables_unchanged": values == item["tables"],
                           "table_sha256_before": table_hash(item["tables"]), "table_sha256_after": table_hash(values),
                           "images_count_unchanged": len(doc.inline_shapes) == item["images"]})
            before_doc = Document(original)
            if "/final/" not in path.as_posix() and "evaluation_tasks_report" not in path.name:
                before = [p.text for p in before_doc.paragraphs]
                after = [p.text for p in doc.paragraphs]
                edited_paragraphs[name] = sum(a != b for a, b in zip(before, after, strict=True))
                code_preserved = True
                for a, b in zip(before_doc.paragraphs, doc.paragraphs, strict=True):
                    if any(r.font.name in ("Consolas", "Courier New") for r in a.runs):
                        code_preserved = code_preserved and a.text == b.text
                checks[-1]["executable_examples_unchanged"] = code_preserved
            word_changed += sha(path) != item["sha256_before"]
    success = not changed_evidence and all(value for check in checks for key, value in check.items() if key.endswith("_unchanged"))
    report = {"date": "2026-10-28", "scope": "Language edit only; no retraining, new inference, relabeling or statistical changes",
              "all_passed": success, "protected_evidence_files": len(protected), "changed_evidence": changed_evidence,
              "edited_markdown_files": markdown_changed, "edited_word_files": word_changed,
              "legacy_paragraph_edits": edited_paragraphs, "document_checks": checks,
              "backup": BACKUP.relative_to(ROOT).as_posix(),
              "limits": "Language revised for readability; original experiment dates, AI-use disclosures and pending independent review retained."}
    output = REPORTS / "maintenance/editorial_verification.json"
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: report[key] for key in ("all_passed", "protected_evidence_files", "changed_evidence", "edited_markdown_files", "edited_word_files", "legacy_paragraph_edits")}, ensure_ascii=False, indent=2))
    if not success:
        print(json.dumps([c for c in checks if any(not v for k, v in c.items() if k.endswith("_unchanged"))], ensure_ascii=False, indent=2))
        raise ValueError("Editorial integrity check failed")


if __name__ == "__main__":
    run()
